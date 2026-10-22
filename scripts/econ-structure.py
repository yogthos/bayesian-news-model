"""Finance and manufacturing value-added shares: the I-O spine against the
national accounts (bead bnm-3ka.1, Econ 0; docs/ECONOMIC-CHARACTER.md 7a,
docs/STOCK-FLOW.md "Finance rows").

Shares are of total value added (the spine's 21 sectors summed; BEA's GDP,
whose industry value added sums to it; OECD's and EU KLEMS's total gross
value added at basic prices), current prices, so units cancel.

  fin  the spine's FIN (LR-WIOD J, WIOD 2016 K64-K66, ICIO K); BEA's
       finance and insurance (NAICS 52, line 55); ISIC4/NACE K elsewhere.
  man  the spine's twelve manufacturing sectors (FOOD ... OTHMAN, REF; their
       PAPER holds publishing, ISIC3 22 / ISIC4 J58); BEA's manufacturing
       (NAICS 31-33, line 12); ISIC4/NACE C elsewhere.

Sources (rows by source):
  lrwiod, wiod16, icio2023  data/derived/io/spine-sectors.csv (scripts/
       io-spine2.py); lrwiod also in 2000 from lrwiod-sectors.csv, the
       year it overlaps WIOD 2016, for chaining
  bea       data/raw/econ/bea-gdp-by-industry: the NAICS historical file
       (1947-1996, November 2021 vintage) and the current release
       (1997-2025, September 2026 vintage)
  oecd      OECD Table 6 (DF_TABLE6, B1G, current prices), K, C and _T
  euklems   EU KLEMS & INTANProd 2023, national_accounts.csv VA_CP, K, C
       and TOT

Writes data/derived/econ/structure-inputs.csv (iso,year,source,fin,man), the
input of bnm.econ-structure (tools/bnm/econ_structure_run.clj writes the
stock), and data/derived/econ/structure-ratios.csv
(iso,segment,benchmark,variable,first,last,n,ratio,min,max: the spine
segment's share over the benchmark's, pooled over the years both cover),
and prints the comparison tables.

Run with the scratchpad venv's python -I (openpyxl, pandas).
"""
import csv
import gzip
import io
import os
import sys
import zipfile

import openpyxl
import pandas as pd

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IO = os.path.join(ROOT, 'data/derived/io')
ECON = os.path.join(ROOT, 'data/raw/econ')
OUT = os.path.join(ROOT, 'data/derived/econ')
MAN = ['FOOD', 'TEX', 'PAPER', 'REF', 'CHEM', 'RUB', 'MINR', 'METAL', 'MACH', 'ELEC', 'TRANSEQ', 'OTHMAN']
BENCH = ['bea', 'oecd', 'euklems']
KLEMS_ISO = {'AT': 'AUT', 'BE': 'BEL', 'BG': 'BGR', 'CY': 'CYP', 'CZ': 'CZE', 'DE': 'DEU', 'DK': 'DNK',
             'EE': 'EST', 'EL': 'GRC', 'ES': 'ESP', 'FI': 'FIN', 'FR': 'FRA', 'HR': 'HRV', 'HU': 'HUN',
             'IE': 'IRL', 'IT': 'ITA', 'JP': 'JPN', 'LT': 'LTU', 'LU': 'LUX', 'LV': 'LVA', 'MT': 'MLT',
             'NL': 'NLD', 'PL': 'POL', 'PT': 'PRT', 'RO': 'ROU', 'SE': 'SWE', 'SI': 'SVN', 'SK': 'SVK',
             'UK': 'GBR', 'US': 'USA'}


def spine():
    """Rows (iso, year, source, fin, man) from the spine, plus LR-WIOD's 2000."""
    sp = pd.read_csv(os.path.join(IO, 'spine-sectors.csv'))
    tot = sp.groupby(['country', 'year', 'source']).va.sum()
    fin = sp[sp.sector == 'FIN'].groupby(['country', 'year', 'source']).va.sum()
    man = sp[sp.sector.isin(MAN)].groupby(['country', 'year', 'source']).va.sum()
    d = pd.DataFrame({'fin': fin / tot, 'man': man / tot}).reset_index()
    lr = pd.read_csv(os.path.join(IO, 'lrwiod-sectors.csv'))
    lr = lr[(lr.year == 2000) & ~lr.country.isin(['ROW', 'xROW'])]
    t = lr.groupby('country').va.sum()
    x = pd.DataFrame({'fin': lr[lr.sector == 'J'].groupby('country').va.sum() / t,
                      'man': lr[lr.sector.str.startswith('D')].groupby('country').va.sum() / t}).reset_index()
    x['year'], x['source'] = 2000, 'lrwiod'
    d = pd.concat([d, x], ignore_index=True)
    return d.rename(columns={'country': 'iso'})[['iso', 'year', 'source', 'fin', 'man']]


def bea_sheet(rows, header_row):
    head = [str(h).strip() if h is not None else '' for h in rows[header_row]]
    by_line = {}
    for r in rows[header_row + 1:]:
        if r[0] is not None and str(r[0]).strip().isdigit():
            by_line[int(str(r[0]).strip())] = r
    out = {}
    for i, h in enumerate(head):
        if h.isdigit():
            vals = [by_line[k][i] for k in (1, 12, 55)]
            if all(isinstance(v, (int, float)) for v in vals):
                out[int(h)] = vals
    return out


def bea():
    """BEA value added shares: GDP (line 1), manufacturing (12), finance and
    insurance (55); historical file before 1997, current release from 1997."""
    z = zipfile.ZipFile(os.path.join(ECON, 'bea-gdp-by-industry', 'AllTablesHist.zip'))
    wb = openpyxl.load_workbook(io.BytesIO(z.read('GDPbyInd_VA_1947-1997.xlsx')), read_only=True)
    rows = list(wb['VA'].iter_rows(values_only=True))
    hist = bea_sheet(rows, next(i for i, r in enumerate(rows) if r[0] == 'Line'))
    z = zipfile.ZipFile(os.path.join(ECON, 'bea-gdp-by-industry', 'GDPbyInd.zip'))
    wb = openpyxl.load_workbook(io.BytesIO(z.read('ValueAdded.xlsx')), read_only=True)
    rows = list(wb['TVA105-A'].iter_rows(values_only=True))
    cur = bea_sheet(rows, next(i for i, r in enumerate(rows) if r[0] == 'Line'))
    out = []
    for y, (gdp, man, fin) in sorted({**{y: v for y, v in hist.items() if y < 1997}, **cur}.items()):
        out.append(('USA', y, 'bea', fin / gdp, man / gdp))
    v97 = [hist[1997][2] / hist[1997][0], cur[1997][2] / cur[1997][0]]
    print('BEA 1997 finance share, 2021 historical vintage %.4f, 2026 current %.4f' % tuple(v97), file=sys.stderr)
    return pd.DataFrame(out, columns=['iso', 'year', 'source', 'fin', 'man'])


def oecd():
    path = os.path.join(ECON, 'oecd-sna-sector-accounts', 'DF_TABLE6_all.csv.gz')
    head = pd.read_csv(path, nrows=0).columns
    keep = ['REF_AREA', 'ACTIVITY', 'TRANSACTION', 'PRICE_BASE', 'UNIT_MEASURE', 'TIME_PERIOD', 'OBS_VALUE']
    t = pd.read_csv(path, usecols=[c for c in head if c.split(':')[0] in keep], dtype=str)
    t.columns = [c.split(':')[0] for c in t.columns]
    for c in ['REF_AREA', 'ACTIVITY', 'TRANSACTION', 'PRICE_BASE', 'UNIT_MEASURE']:
        t[c] = t[c].str.split(':').str[0]
    t = t[(t.TRANSACTION == 'B1G') & (t.PRICE_BASE == 'V') & (t.UNIT_MEASURE == 'XDC')
          & t.ACTIVITY.isin(['K', 'C', '_T'])]
    t['OBS_VALUE'] = pd.to_numeric(t.OBS_VALUE, errors='coerce')
    p = t.pivot_table(index=['REF_AREA', 'TIME_PERIOD'], columns='ACTIVITY', values='OBS_VALUE',
                      aggfunc='first').dropna()
    p = p[p['_T'] > 0].reset_index()
    return pd.DataFrame({'iso': p.REF_AREA, 'year': p.TIME_PERIOD.astype(int), 'source': 'oecd',
                         'fin': p.K / p['_T'], 'man': p.C / p['_T']})


def euklems():
    k = pd.read_csv(os.path.join(ECON, 'euklems-intanprod-2023', 'national_accounts.csv'),
                    usecols=['nace_r2_code', 'geo_code', 'year', 'VA_CP'])
    k = k[k.nace_r2_code.isin(['K', 'C', 'TOT']) & k.geo_code.isin(KLEMS_ISO)]
    p = k.pivot_table(index=['geo_code', 'year'], columns='nace_r2_code', values='VA_CP', aggfunc='first').dropna()
    p = p[p.TOT > 0].reset_index()
    return pd.DataFrame({'iso': p.geo_code.map(KLEMS_ISO), 'year': p.year.astype(int), 'source': 'euklems',
                         'fin': p.K / p.TOT, 'man': p.C / p.TOT})


def ratios(d):
    """Each spine segment's share over each benchmark's, pooled (sum over
    sum) over the years both cover."""
    sp = d[~d.source.isin(BENCH)]
    out = []
    for b in BENCH:
        m = sp.merge(d[d.source == b], on=['iso', 'year'], suffixes=('', '_b'))
        for (iso, seg), g in m.groupby(['iso', 'source']):
            for v in ('fin', 'man'):
                g2 = g[(g[v] > 0) & (g[v + '_b'] > 0)]
                if len(g2):
                    r = g2[v] / g2[v + '_b']
                    out.append((iso, seg, b, v, g2.year.min(), g2.year.max(), len(g2),
                                g2[v].sum() / g2[v + '_b'].sum(), r.min(), r.max()))
    return pd.DataFrame(out, columns=['iso', 'segment', 'benchmark', 'variable', 'first', 'last', 'n',
                                      'ratio', 'min', 'max'])


def main():
    os.makedirs(OUT, exist_ok=True)
    s = spine()
    countries = set(s.iso)
    b = pd.concat([bea(), oecd(), euklems()], ignore_index=True)
    b = b[b.iso.isin(countries)]
    d = pd.concat([s, b], ignore_index=True).sort_values(['iso', 'year', 'source'])
    d.to_csv(os.path.join(OUT, 'structure-inputs.csv'), index=False, float_format='%.6g')
    r = ratios(d)
    r.to_csv(os.path.join(OUT, 'structure-ratios.csv'), index=False, float_format='%.4g')
    pd.set_option('display.width', 200)
    print('rows', len(d), {k: len(g) for k, g in d.groupby('source')}, file=sys.stderr)
    us = d[d.iso == 'USA'].pivot_table(index='year', columns='source', values=['fin', 'man'])
    yrs = [1965, 1970, 1975, 1980, 1985, 1990, 1995, 1999, 2000, 2005, 2010, 2014, 2015, 2020]
    for v in ('fin', 'man'):
        u = us[v].reindex(yrs)
        u['spine'] = u[['lrwiod', 'wiod16', 'icio2023']].bfill(axis=1).iloc[:, 0]
        u.loc[2000, 'spine'] = u.loc[2000, 'wiod16']
        u['spine/bea'] = u.spine / u.bea
        print('\nUSA %s share, %%; spine/BEA' % v)
        print((u[['lrwiod', 'wiod16', 'icio2023', 'bea', 'oecd', 'euklems']] * 100).round(2)
              .join(u['spine/bea'].round(3)).to_string())
    print('\nPooled ratios, spine segment over benchmark (fin, then man):')
    for v in ('fin', 'man'):
        t = r[r.variable == v].pivot_table(index='iso', columns=['segment', 'benchmark'], values='ratio')
        print('\n' + v)
        print(t.round(2).to_string())


if __name__ == '__main__':
    main()
